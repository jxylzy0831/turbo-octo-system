from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[1]
source = root / 'drafts/废土电影情感本_完整创作提示词_v1.md'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(2)
sec.left_margin = sec.right_margin = Cm(2.2)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'List Bullet', 'List Number']:
    s = doc.styles[name]
    s.font.name = 'Microsoft YaHei'
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    s.font.size = Pt(10.5)
    s.paragraph_format.line_spacing = 1.15
    s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.widow_control = True
for name, size in [('Title', 20), ('Heading 1', 14), ('Heading 2', 11.5)]:
    s = doc.styles[name]
    s.font.size = Pt(size)
    s.font.bold = True
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.space_before = Pt(10)

lines = source.read_text(encoding='utf-8').splitlines()
i = 0
while i < len(lines):
    line = lines[i].strip()
    i += 1
    if not line:
        continue
    if line.startswith('|'):
        rows = [line]
        while i < len(lines) and lines[i].strip().startswith('|'):
            rows.append(lines[i].strip())
            i += 1
        rows = [r for r in rows if not re.match(r'^\|[\s|:-]+\|$', r)]
        table = doc.add_table(rows=0, cols=3)
        table.autofit = False
        for col, width in zip(table.columns, [3.5, 3, 10.1]):
            col.width = Cm(width)
        for idx, row in enumerate(rows):
            cells = table.add_row().cells
            trpr = table.rows[-1]._tr.get_or_add_trPr()
            trpr.append(OxmlElement('w:cantSplit'))
            if idx == 0:
                trpr.append(OxmlElement('w:tblHeader'))
            for cell, value, width in zip(cells, row.strip('|').split('|'), [3.5, 3, 10.1]):
                cell.width = Cm(width)
                cell.text = value.strip()
                cell.vertical_alignment = 1
                pr = cell._tc.get_or_add_tcPr()
                borders = OxmlElement('w:tcBorders')
                for edge in ['top','left','bottom','right']:
                    el = OxmlElement('w:' + edge)
                    for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:
                        el.set(qn('w:'+k),v)
                    borders.append(el)
                pr.append(borders)
                margins = OxmlElement('w:tcMar')
                for edge in ['top','left','bottom','right']:
                    el = OxmlElement('w:'+edge)
                    el.set(qn('w:w'),'90')
                    el.set(qn('w:type'),'dxa')
                    margins.append(el)
                pr.append(margins)
                if idx == 0:
                    shade = OxmlElement('w:shd')
                    shade.set(qn('w:fill'),'E8EDF2')
                    pr.append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(2)
                    for run in p.runs:
                        run.font.size = Pt(9)
                        run.bold = idx == 0
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
    elif line.startswith('# '):
        doc.add_paragraph(line[2:], 'Title')
    elif line.startswith('## '):
        doc.add_paragraph(line[3:].replace('、', ' '), 'Heading 1')
    elif line.startswith('### '):
        doc.add_paragraph(line[4:], 'Heading 2')
    elif line.startswith('- '):
        doc.add_paragraph(line[2:], 'List Bullet')
    elif re.match(r'^\d+\. ', line):
        doc.add_paragraph(re.sub(r'^\d+\. ', '', line), 'List Number')
    else:
        doc.add_paragraph(line)
target = root / 'drafts/废土电影情感本_完整创作提示词_v1.docx'
for el in list(doc.styles.element.iter(qn('w:pBdr'))) + list(doc.element.iter(qn('w:pBdr'))):
    el.getparent().remove(el)
doc.save(target)
print(target)
