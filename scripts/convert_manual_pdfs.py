from __future__ import annotations

import json
import re
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "source" / "originals" / "非《流氓叙事》的组织者手册参考"
TEMP = ROOT / "临时存储" / "pdf_to_docx"
OUTPUT = TEMP / "converted_manuals"
OCR_DIR = TEMP / "ocr_jobs"
MANIFEST = ROOT / "analysis" / "外部组织者手册_SHA256_初始清单.json"

SKIP_PAGES = {
    "《青白》组织者手册.pdf": {1},
    "暗夜降至手册（最终版）.pdf": {1},
    "欲缚.pdf": {1},
    "碌碌无为手册.pdf": {1},
    "绿洲手册2.0.pdf": {1},
    "窃云台组织者手册.pdf": {1},
    "草芥手册.pdf": {1},
    "醉凌云手册.pdf": {1},
}
OCR_FILES = {
    "欲缚.pdf": OCR_DIR / "job1.jsonl",
    "碌碌无为手册.pdf": OCR_DIR / "job2.jsonl",
}
PAGE_NUMBER = re.compile(
    r"^(?:(?:第\s*)?\d{1,4}(?:\s*页)?(?:\s*(?:共|/)\s*\d{1,4}\s*页?)?|\d{1,4}\s*/\s*\d{1,4})$"
)
SUSPECT = re.compile(r"[\ufffd�□■]{1,}|\?{2,}")


def ocr_pages(pdf_name: str) -> dict[int, dict]:
    path = OCR_FILES.get(pdf_name)
    if not path or not path.exists():
        return {}
    pages = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        row = json.loads(line)
        pages[int(row["page"])] = row
    return pages


def page_lines(text: str) -> list[str]:
    return [line.rstrip() for line in (text or "").replace("\x00", "").splitlines()]


def repeated_marginal_lines(lines_by_page: dict[int, list[str]]) -> set[str]:
    occurrences: Counter[str] = Counter()
    for lines in lines_by_page.values():
        clean = [(line.strip()) for line in lines if line.strip()]
        if not clean:
            continue
        for line in set(clean[:2] + clean[-2:]):
            if not PAGE_NUMBER.fullmatch(line):
                occurrences[line] += 1
    minimum = max(4, int(len(lines_by_page) * 0.12))
    return {line for line, count in occurrences.items() if count >= minimum and len(line) <= 60}


def configure_document(doc: Document, width_pt: float, height_pt: float) -> None:
    section = doc.sections[0]
    section.page_width = Pt(width_pt)
    section.page_height = Pt(height_pt)
    margin = min(30.0, width_pt * 0.055, height_pt * 0.045)
    section.left_margin = Pt(margin)
    section.right_margin = Pt(margin)
    section.top_margin = Pt(margin)
    section.bottom_margin = Pt(margin)
    section.header_distance = Pt(0)
    section.footer_distance = Pt(0)
    normal = doc.styles["Normal"]
    normal.font.name = "NSimSun"
    normal.font.size = Pt(8)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "新宋体")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0


def write_page(doc: Document, lines: list[str], is_ocr: bool) -> int:
    suspect_count = 0
    if not any(line.strip() for line in lines):
        return 0
    for line in lines:
        if not line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(3)
            continue
        if PAGE_NUMBER.fullmatch(line.strip()):
            continue
        p = doc.add_paragraph()
        lead = len(line) - len(line.lstrip(" "))
        p.paragraph_format.left_indent = Pt(min(lead * (4.1 if is_ocr else 3.7), 240))
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(line.lstrip(" "))
        run.font.name = "NSimSun"
        run.font.size = Pt(9 if is_ocr else 8)
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "新宋体")
        if SUSPECT.search(line):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            suspect_count += 1
    return suspect_count


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    hashes = {row["file"]: row for row in json.loads(MANIFEST.read_text(encoding="utf-8-sig"))}
    report = []
    for pdf_path in sorted(SOURCE.glob("*.pdf")):
        reader = PdfReader(str(pdf_path))
        ocr = ocr_pages(pdf_path.name)
        extracted: dict[int, list[str]] = {}
        page_sizes: dict[int, tuple[float, float]] = {}
        page_sources: dict[int, str] = {}
        for number, page in enumerate(reader.pages, start=1):
            box = page.mediabox
            page_sizes[number] = (float(box.width), float(box.height))
            row = ocr.get(number)
            if row is not None:
                lines = [str(line.get("text", "")) for line in row.get("lines", [])]
                page_sources[number] = "Windows OCR"
            else:
                try:
                    text = page.extract_text(extraction_mode="layout") or ""
                except Exception:
                    text = page.extract_text() or ""
                lines = page_lines(text)
                page_sources[number] = "PDF text"
            extracted[number] = lines

        repeated = repeated_marginal_lines(extracted)
        output_path = OUTPUT / f"{pdf_path.stem}_去封面正文可编辑.docx"
        doc = Document()
        width, height = page_sizes[1]
        configure_document(doc, width, height)
        included = [n for n in range(1, len(reader.pages) + 1) if n not in SKIP_PAGES.get(pdf_path.name, set())]
        suspects = 0
        empty = 0
        for index, number in enumerate(included):
            lines = [line for line in extracted[number] if line.strip() not in repeated]
            if not any(line.strip() for line in lines):
                empty += 1
            suspects += write_page(doc, lines, page_sources[number] == "Windows OCR")
            if index < len(included) - 1:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.core_properties.title = pdf_path.stem + " 正文可编辑版"
        doc.core_properties.subject = "由原 PDF 逐页转换的可编辑正文"
        doc.save(output_path)
        report.append(
            {
                "source": pdf_path.name,
                "sha256": hashes[pdf_path.name]["sha256"],
                "source_pages": len(reader.pages),
                "removed_cover_pages": sorted(SKIP_PAGES.get(pdf_path.name, set())),
                "word_pages_expected": len(included),
                "ocr_pages": sum(v == "Windows OCR" for v in page_sources.values()),
                "empty_pages_preserved": empty,
                "uncertain_lines_highlighted": suspects,
                "output": str(output_path.relative_to(ROOT)),
                "bytes": output_path.stat().st_size,
            }
        )
        print(f"{pdf_path.name}: {len(included)} pages -> {output_path.name}", flush=True)
    (ROOT / "analysis" / "外部组织者手册_Word转换清单.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )


if __name__ == "__main__":
    main()
