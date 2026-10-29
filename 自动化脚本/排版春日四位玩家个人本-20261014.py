"""将四位玩家最新版 Markdown 排为独立 Word，保留全部故事段落。"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / '成品/春日四位玩家个人本-20261014-排版第1版'
QA = ROOT / '质检/春日四位玩家个人本-20261014-排版第1版'
NAMES = ('恩书', '知夏', '多允', '瑞真')
PARAGRAPH_SPACING = {
    ('知夏', '第二幕 开学派对'): 3,
    ('知夏', '第四幕 真相与清晨会'): 4,
    ('瑞真', '第六幕 奉还礼'): 2,
}


def latest_source(name: str) -> tuple[Path, int]:
    pattern = re.compile(rf'^{name}-人物本-(\d{{8}})-(?:重写)?第(\d+)版\.md$')
    choices = []
    for path in (ROOT / '草稿').glob(f'{name}-人物本-*.md'):
        match = pattern.fullmatch(path.name)
        if match:
            choices.append((int(match[1]), int(match[2]), path.stat().st_mtime_ns, path))
    if not choices:
        raise FileNotFoundError(f'找不到{name}的玩家本')
    _, version, _, path = max(choices)
    return path, version


def set_font(style, east: str, size: float, bold: bool = False):
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attribute, value in [('eastAsia', east), ('ascii', 'Times New Roman'),
                             ('hAnsi', 'Times New Roman'), ('cs', 'Times New Roman')]:
        fonts.set(qn('w:' + attribute), value)
    for attribute in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        fonts.attrib.pop(qn('w:' + attribute), None)
    lang = style.element.get_or_add_rPr().find(qn('w:lang'))
    if lang is None:
        lang = OxmlElement('w:lang')
        style.element.get_or_add_rPr().append(lang)
    lang.set(qn('w:val'), 'en-US')
    lang.set(qn('w:eastAsia'), 'zh-CN')


def page_field(paragraph):
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    run = OxmlElement('w:r')
    text = OxmlElement('w:t')
    text.text = '1'
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def blocks_from_markdown(text: str):
    blocks = []
    pending = []

    def flush():
        if pending:
            blocks.append(('body', '\n'.join(pending)))
            pending.clear()

    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            flush()
        elif line.startswith('# '):
            flush()
        elif line.startswith('## '):
            flush()
            blocks.append(('act', line[3:]))
        elif line.startswith('### '):
            flush()
            blocks.append(('ending', line[4:]))
        else:
            pending.append(raw_line)
    flush()
    return blocks


def build(source: Path, name: str, version: int, overwrite: bool):
    target = OUTPUT / f'春日-{name}-人物本-第{version}版.docx'
    if target.exists() and not overwrite:
        raise FileExistsError(target)
    source_bytes = source.read_bytes()
    source_text = source.read_text(encoding='utf-8-sig')
    blocks = blocks_from_markdown(source_text)
    acts = [text for kind, text in blocks if kind == 'act']
    assert len(acts) == 7 and acts[1] == '第二幕 开学派对'
    endings = [text for kind, text in blocks if kind == 'ending']
    assert endings == ([] if name == '恩书' else ['生还', '未获救'])

    document = Document()
    for border in document.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    normal = document.styles['Normal']
    set_font(normal, '宋体', 12)
    paragraph = normal.paragraph_format
    paragraph.first_line_indent = Pt(24)
    paragraph.line_spacing = Pt(21)
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(5)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.widow_control = True
    paragraph.keep_with_next = False
    paragraph.keep_together = False

    for style_name, east, size, bold, after in [
        ('Title', '宋体', 38, False, 20),
        ('Subtitle', '微软雅黑', 16, False, 0),
        ('Heading 1', '微软雅黑', 18, True, 19),
        ('Heading 2', '微软雅黑', 14, True, 12),
        ('Header', '宋体', 9, False, 0),
        ('Footer', '宋体', 9, False, 0),
    ]:
        style = document.styles[style_name]
        set_font(style, east, size, bold)
        paragraph = style.paragraph_format
        paragraph.first_line_indent = Pt(0)
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(after)
        paragraph.line_spacing = 1.2
        paragraph.keep_with_next = style_name in ['Title', 'Heading 1', 'Heading 2']
        paragraph.keep_together = True
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    cover = document.sections[0]
    cover.page_width, cover.page_height = Mm(210), Mm(297)
    cover.top_margin, cover.bottom_margin = Mm(25), Mm(23)
    cover.left_margin, cover.right_margin = Mm(26), Mm(26)
    cover.header_distance, cover.footer_distance = Mm(12), Mm(11)

    title = document.add_paragraph('春日', 'Title')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(200)
    subtitle = document.add_paragraph(f'{name}人物本', 'Subtitle')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    story = document.add_section(WD_SECTION_START.NEW_PAGE)
    story.header.is_linked_to_previous = False
    story.footer.is_linked_to_previous = False
    numbering = OxmlElement('w:pgNumType')
    numbering.set(qn('w:start'), '1')
    story._sectPr.append(numbering)
    header = story.header.paragraphs[0]
    header.style = document.styles['Header']
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.add_run(f'春日  {name}')
    footer = story.footer.paragraphs[0]
    footer.style = document.styles['Footer']
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page_field(footer)

    actual = []
    first_act = True
    active_act = ''
    for kind, text in blocks:
        if kind == 'act':
            active_act = text
            paragraph = document.add_paragraph(text, 'Heading 1')
            paragraph.paragraph_format.page_break_before = not first_act
            first_act = False
        elif kind == 'ending':
            paragraph = document.add_paragraph(text, 'Heading 2')
            paragraph.paragraph_format.page_break_before = text == '未获救'
        else:
            paragraph = document.add_paragraph(text)
            if (name, active_act) in PARAGRAPH_SPACING:
                paragraph.paragraph_format.space_after = Pt(PARAGRAPH_SPACING[name, active_act])
            if len(text) <= 35:
                paragraph.paragraph_format.keep_together = True
        actual.append(paragraph.text)
    expected = [text for _, text in blocks]
    assert actual == expected

    properties = document.core_properties
    properties.title = f'春日 {name}人物本'
    properties.subject = '人物本'
    properties.author = ''
    properties.last_modified_by = ''
    properties.comments = ''
    properties.keywords = ''
    properties.language = 'zh-CN'
    for setting_name in ['updateFields', 'doNotTrackMoves', 'doNotTrackFormatting']:
        element = OxmlElement('w:' + setting_name)
        element.set(qn('w:val'), 'true')
        document.settings.element.append(element)
    document.save(target)

    saved = Document(target)
    saved_story = [paragraph.text for paragraph in saved.paragraphs if paragraph.text][2:]
    assert saved_story == expected, f'{name}转换后段落不一致'
    assert hashlib.sha256(source.read_bytes()).digest() == hashlib.sha256(source_bytes).digest()
    return {
        '人物': name,
        '版本': version,
        '来源': str(source.relative_to(ROOT)),
        '来源sha256': hashlib.sha256(source_bytes).hexdigest(),
        '文档': str(target.relative_to(ROOT)),
        '文档sha256': hashlib.sha256(target.read_bytes()).hexdigest(),
        '故事段落与标题数': len(expected),
        '故事汉字': len(re.findall(r'[\u4e00-\u9fff]', ''.join(expected))),
        '分幕': acts,
        '结尾': endings,
        '逐段一致': True,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--重排', action='store_true', help='只用于本轮渲染后的版式返修')
    parser.add_argument('--人物', nargs='+', choices=NAMES, default=list(NAMES))
    args = parser.parse_args()
    sources = [(name, *latest_source(name)) for name in args.人物]
    OUTPUT.mkdir(parents=True, exist_ok=True)
    QA.mkdir(parents=True, exist_ok=True)
    manifest = QA / '正文与来源核对.json'
    previous = {item['人物']: item for item in json.loads(manifest.read_text(encoding='utf-8'))} if manifest.exists() else {}
    for name, source, version in sources:
        record = build(source, name, version, args.重排)
        previous[name] = record
        print(f'{name} 第{version}版 {record["故事段落与标题数"]}段 {record["文档"]}')
    records = [previous[name] for name in NAMES if name in previous]
    manifest.write_text(json.dumps(records, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


if __name__ == '__main__':
    main()
