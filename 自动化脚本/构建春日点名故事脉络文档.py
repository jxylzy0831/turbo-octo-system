from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / '草稿/春日点名-完整故事脉络-20261009-第9版.md'
QA = ROOT / '质检/春日点名-第9版Word排版'
OUTPUT = QA / '春日点名-完整故事脉络-第9版-对外阅读版.docx'


def set_font(style, east='宋体', west='Arial', size=12, bold=False):
    style.font.name = west
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key, value in [('eastAsia', east), ('ascii', west), ('hAnsi', west), ('cs', west)]:
        fonts.set(qn('w:' + key), value)
    for key in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        fonts.attrib.pop(qn('w:' + key), None)


def field(paragraph, instruction):
    run = paragraph.add_run()
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), instruction)
    run._r.addnext(fld)


def table_border(table):
    properties = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        element = OxmlElement('w:' + name)
        for key, value in [('val', 'single'), ('sz', '5'), ('color', 'D9D9D9')]:
            element.set(qn('w:' + key), value)
        borders.append(element)
    properties.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for name, value in [('top', 125), ('bottom', 125), ('left', 130), ('right', 130)]:
        element = OxmlElement('w:' + name)
        element.set(qn('w:w'), str(value))
        element.set(qn('w:type'), 'dxa')
        margins.append(element)
    properties.append(margins)


def main():
    QA.mkdir(parents=True, exist_ok=True)
    source = SOURCE.read_text(encoding='utf-8-sig')
    lines = source.splitlines()
    doc = Document()
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin, section.bottom_margin = Mm(22), Mm(22)
    section.left_margin, section.right_margin = Mm(23), Mm(23)
    section.header_distance, section.footer_distance = Mm(10), Mm(11)
    section.different_first_page_header_footer = True

    normal = doc.styles['Normal']
    set_font(normal)
    normal.paragraph_format.line_spacing = Pt(21)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Pt(24)
    for name, size, before, after in [('Title', 27, 0, 7), ('Subtitle', 12, 0, 14),
                                       ('Heading 1', 16, 16, 8), ('Heading 2', 12, 10, 5)]:
        style = doc.styles[name]
        set_font(style, east='微软雅黑', size=size, bold=(name != 'Subtitle'))
        fmt = style.paragraph_format
        fmt.first_line_indent = Pt(0)
        fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt.space_before, fmt.space_after = Pt(before), Pt(after)
        fmt.line_spacing = 1.2
        fmt.keep_with_next = True
        fmt.keep_together = True
    for name in ['Header', 'Footer']:
        style = doc.styles[name]
        set_font(style, east='微软雅黑', size=9)
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = 1

    doc.core_properties.title = '春日点名 完整故事脉络 第9版'
    doc.core_properties.subject = '人物与专属恋陪对应及七幕故事脉络'
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.core_properties.comments = ''
    update = OxmlElement('w:updateFields')
    update.set(qn('w:val'), 'true')
    doc.settings.element.append(update)

    title = doc.add_paragraph('春日点名', 'Title')
    title.paragraph_format.space_before = Pt(3)
    doc.add_paragraph('完整故事脉络 第9版', 'Subtitle')

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.add_run('春日点名 完整故事脉络')
    for footer in [section.footer, section.first_page_footer]:
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run('第 ')
        field(p, 'PAGE')
        p.add_run(' 页  共 ')
        field(p, 'NUMPAGES')
        p.add_run(' 页')

    expected = []
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        index += 1
        if not line or line.startswith('# '):
            continue
        if line.startswith('## '):
            text = line[3:]
            heading = doc.add_paragraph(text, 'Heading 1')
            if text == '人物与代号规则':
                heading.paragraph_format.space_before = Pt(10)
            if text == '第一幕':
                heading.paragraph_format.page_break_before = True
                heading.paragraph_format.space_before = Pt(0)
            expected.append(text)
        elif line.startswith('### '):
            text = line[4:]
            doc.add_paragraph(text, 'Heading 2')
            expected.append(text)
        elif line.startswith('|'):
            rows = [line]
            while index < len(lines) and lines[index].strip().startswith('|'):
                rows.append(lines[index].strip())
                index += 1
            rows = [row for row in rows if not re.fullmatch(r'[|\s:-]+', row)]
            values = [[cell.strip() for cell in row.strip('|').split('|')] for row in rows]
            table = doc.add_table(rows=0, cols=4)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            widths = [22, 38, 35, 69]
            for column, width in zip(table.columns, widths):
                column.width = Mm(width)
            table_border(table)
            for row_no, value_row in enumerate(values):
                cells = table.add_row().cells
                trpr = cells[0]._tc.getparent().get_or_add_trPr()
                trpr.append(OxmlElement('w:cantSplit'))
                if row_no == 0:
                    trpr.append(OxmlElement('w:tblHeader'))
                for column_no, text in enumerate(value_row):
                    cell = cells[column_no]
                    cell.width = Mm(widths[column_no])
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    shading = OxmlElement('w:shd')
                    shading.set(qn('w:fill'), '293A4C' if row_no == 0 else ('F2F5F7' if row_no % 2 else 'FFFFFF'))
                    cell._tc.get_or_add_tcPr().append(shading)
                    paragraph = cell.paragraphs[0]
                    paragraph.paragraph_format.first_line_indent = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = Pt(17)
                    paragraph.paragraph_format.keep_with_next = row_no < len(values) - 1
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if column_no in (0, 2) else WD_ALIGN_PARAGRAPH.LEFT
                    run = paragraph.add_run(text)
                    run.font.name = 'Arial'
                    run._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '微软雅黑')
                    run.font.size = Pt(10.5)
                    run.font.bold = row_no == 0 or column_no in (0, 2)
                    run.font.color.rgb = RGBColor.from_string('FFFFFF' if row_no == 0 else '000000')
                    expected.append(text)
            gap = doc.add_paragraph()
            gap.paragraph_format.first_line_indent = Pt(0)
            gap.paragraph_format.space_after = Pt(1)
            gap.paragraph_format.line_spacing = Pt(4)
            gap.add_run().font.size = Pt(4)
        else:
            doc.add_paragraph(line)
            expected.append(line)

    doc.save(OUTPUT)
    # Check all source blocks against saved OOXML, including table cells in order.
    restored = Document(OUTPUT)
    actual = []
    for node in restored.element.body:
        if node.tag == qn('w:p'):
            text = ''.join(node.xpath('.//w:t/text()'))
            if text and text not in ['春日点名', '完整故事脉络 第9版']:
                actual.append(text)
        elif node.tag == qn('w:tbl'):
            for cell in node.xpath('./w:tr/w:tc'):
                actual.append(''.join(cell.xpath('.//w:t/text()')))
    if actual != expected:
        raise RuntimeError('正文与第9版内容不一致')
    report = {'source': str(SOURCE), 'source_sha256': hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
              'working_docx': str(OUTPUT), 'content_blocks': len(expected), 'content_match': True,
              'page_size': 'A4', 'body_font': '宋体 12 pt', 'line_spacing': '21 pt'}
    (QA / '内容核对.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False))


if __name__ == '__main__':
    main()
