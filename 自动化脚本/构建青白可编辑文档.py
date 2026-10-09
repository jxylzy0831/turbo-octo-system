from __future__ import annotations
import json, re, statistics, hashlib
from collections import Counter
from pathlib import Path
import pdfplumber
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_BREAK
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
PDF=ROOT/'原始资料/原始文件/非《流氓叙事》的组织者手册参考/《青白》组织者手册.pdf'
OCR=ROOT/'临时存储/pdf_to_docx/qingbai_raster_ocr.jsonl'
OUT=ROOT/'成品/《青白》组织者手册_可编辑文字识别正文版.docx'
REPORT=ROOT/'分析/青白可编辑文字识别转换清单.json'

# Windows OCR inserts spaces between Chinese glyphs. Remove those while retaining
# useful separation between adjacent Latin words and digits.
def normalize(s:str)->str:
    s=s.replace('\u3000',' ')
    s=re.sub(r'(?<=[\u3400-\u9fff])\s+|\s+(?=[\u3400-\u9fff])','',s)
    return re.sub(r'\s+',' ',s).strip()

def rgb(color):
    if color is None:return None
    if isinstance(color,(int,float)): vals=[float(color)]
    elif isinstance(color,(list,tuple)): vals=[float(v) for v in color]
    else:return None
    if len(vals)==1: vals*=3
    if len(vals)<3:return None
    return tuple(max(0,min(255,round(x*255))) for x in vals[:3])

def is_yellow(color):
    c=rgb(color)
    return c is not None and c[0]>220 and c[1]>200 and c[2]<80

def group_rows(chars):
    cs=[c for c in chars if str(c.get('text','')).strip()]
    cs.sort(key=lambda c:(float(c.get('top',0)),float(c.get('x0',0))))
    rows=[]
    for c in cs:
        if rows and abs(float(c.get('top',0))-statistics.median(float(x.get('top',0)) for x in rows[-1]))<=1.25:
            rows[-1].append(c)
        else: rows.append([c])
    for row in rows:
        row.sort(key=lambda c:float(c.get('x0',0)))
    return rows

def row_for_index(i,nlines,nrows):
    if nrows<=1:return 0
    return round(i*(nrows-1)/max(1,nlines-1))

def source_style(char, rects):
    color=rgb(char.get('non_stroking_color'))
    if color is None or max(color)-min(color)<12 and max(color)<55: color=(0,0,0)
    bold=any(k in str(char.get('fontname','')).lower() for k in ('bold','heavy','black','semibold'))
    italic='italic' in str(char.get('fontname','')).lower() or 'oblique' in str(char.get('fontname','')).lower()
    x0,x1=float(char['x0']),float(char['x1']); top=float(char['top']); bottom=float(char['bottom'])
    underline=False; highlight=False
    for r in rects:
        rx0,rx1=float(r.get('x0',0)),float(r.get('x1',0)); rt,rb=float(r.get('top',0)),float(r.get('bottom',0))
        overlap=rx1>x0+0.1 and rx0<x1-0.1
        if not overlap: continue
        if is_yellow(r.get('non_stroking_color')) and rt<=bottom+1 and rb>=top-1: highlight=True
        if rb-rt<1.3 and rt>=bottom-2.0 and rt<=bottom+2.0: underline=True
    return (color,bold,italic,underline,highlight)

def merged_lines(ocr_lines, rows, rects):
    usable=[normalize(str(s)) for s in ocr_lines]
    usable=[s for s in usable if s]
    if not usable:return []
    row_ids=[row_for_index(i,len(usable),len(rows)) for i in range(len(usable))]
    buckets=[]
    for s,ri in zip(usable,row_ids):
        row=rows[ri] if rows else []
        glyphs=[c for c in row if str(c.get('text','')).strip()]
        glyphs.sort(key=lambda c:float(c.get('x0',0)))
        styles=[]
        for j,ch in enumerate(s):
            if glyphs:
                src=glyphs[min(len(glyphs)-1,int(j*len(glyphs)/max(1,len(s))))]
                styles.append(source_style(src,rects))
            else:styles.append(((0,0,0),False,False,False,False))
        left=min([float(c['x0']) for c in row] or [72.0])
        right=max([float(c['x1']) for c in row] or [525.0])
        size=statistics.median([float(c.get('size',10)) for c in row]) if row else 10.0
        top=min([float(c['top']) for c in row] or [70.0])
        if buckets and buckets[-1]['row']==ri:
            buckets[-1]['parts'].append((s,styles))
        else:buckets.append({'row':ri,'parts':[(s,styles)],'left':left,'right':right,'size':size,'top':top})
    out=[]
    for b in buckets:
        chars=[]; st=[]
        for k,(s,ss) in enumerate(b['parts']):
            if k:
                chars.extend('  '); st.extend([((0,0,0),False,False,False,False)]*2)
            chars.extend(s); st.extend(ss)
        b['text']=''.join(chars); b['styles']=st
        out.append(b)
    return out

def add_styled_text(paragraph, text, styles, size):
    if not text:return
    i=0
    while i<len(text):
        style=styles[i] if i<len(styles) else ((0,0,0),False,False,False,False)
        j=i+1
        while j<len(text) and (styles[j] if j<len(styles) else ((0,0,0),False,False,False,False))==style:j+=1
        run=paragraph.add_run(text[i:j])
        run.font.name='宋体'; run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
        run.font.size=Pt(max(7.5,min(18.0,float(size))))
        color,bold,italic,underline,highlight=style
        run.font.color.rgb=RGBColor(*color)
        run.bold=bold;run.italic=italic;run.underline=underline
        if highlight:run.font.highlight_color=WD_COLOR_INDEX.YELLOW
        i=j

def main():
    ocr_by_page={}
    for line in OCR.read_text(encoding='utf-8').splitlines():
        r=json.loads(line);ocr_by_page[int(r['page'])]=r
    if len(ocr_by_page)!=182:raise RuntimeError(f'Expected OCR for 182 body pages, got {len(ocr_by_page)}')
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Pt(595.32);sec.page_height=Pt(841.92)
    sec.left_margin=Pt(0);sec.right_margin=Pt(0);sec.top_margin=Pt(0);sec.bottom_margin=Pt(0)
    sec.header_distance=Pt(0);sec.footer_distance=Pt(0)
    normal=doc.styles['Normal'];normal.font.name='宋体';normal.font.size=Pt(10)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
    normal.paragraph_format.space_after=Pt(0);normal.paragraph_format.space_before=Pt(0)
    normal.paragraph_format.widow_control=False
    report={'source':str(PDF.relative_to(ROOT)),'source_sha256':hashlib.sha256(PDF.read_bytes()).hexdigest().upper(),
            'source_pages':183,'cover_pages_removed':[1],'body_pages':182,
            'ocr_engine':'Windows Media OCR (Simplified Chinese) applied to 144 dpi raster pages',
            'images_preserved':False,'editable_text':True,'output':str(OUT.relative_to(ROOT)),'pages':[]}
    with pdfplumber.open(PDF) as pdf:
        for page_no in range(2,184):
            source=pdf.pages[page_no-1];rows=group_rows(source.chars)
            record=ocr_by_page[page_no]
            if record.get('error'):raise RuntimeError(f'Page {page_no}: {record["error"]}')
            lines=merged_lines(record.get('lines',[]),rows,source.rects)
            nonempty=[x for x in lines if x['text']]
            if not nonempty:raise RuntimeError(f'Page {page_no} produced no OCR text')
            if len(doc.paragraphs):doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            y0=nonempty[0]['top']; y1=nonempty[-1]['top']
            # Allocate the source's vertical span across recognized lines when OCR
            # has split or combined a source row.
            targets=[]
            if len(nonempty)==1:targets=[y0]
            else:targets=[y0+(y1-y0)*i/(len(nonempty)-1) for i in range(len(nonempty))]
            prev_top=0.0
            for i,(b,target) in enumerate(zip(nonempty,targets)):
                p=doc.add_paragraph();pf=p.paragraph_format
                pf.left_indent=Pt(max(36,min(220,b['left'])))
                pf.right_indent=Pt(0)
                # OCR can merge adjacent cell fragments that were split in the
                # source layout. Keep each editable line within the page width.
                width=max(80,595.32-b['left']-36)
                fitted=width/max(1,len(b['text'])*0.94)
                size=max(7.5,min(16,b['size'],fitted))
                line_h=max(size+1,11)
                pf.line_spacing=Pt(line_h)
                pf.space_before=Pt(max(0,target if i==0 else target-prev_top-line_h))
                pf.space_after=Pt(0);pf.widow_control=False
                # Short rows are often headings or centered callouts in the source.
                if len(b['text'])<22 and 130<b['left']<450:
                    p.alignment=1
                add_styled_text(p,b['text'],b['styles'],size)
                prev_top=target
            report['pages'].append({'source_pdf_page':page_no,'word_page':page_no-1,'ocr_lines':len(record.get('lines',[])),
                                    'output_text_lines':len(nonempty),'characters':sum(len(x['text']) for x in nonempty)})
    doc.core_properties.title='《青白》组织者手册 可编辑正文 OCR 版'
    doc.core_properties.subject='源 PDF 第2至183页 OCR 转换为可编辑文字；图片不保留'
    OUT.parent.mkdir(parents=True,exist_ok=True);doc.save(OUT)
    report['output_bytes']=OUT.stat().st_size
    REPORT.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'output':str(OUT),'bytes':report['output_bytes'],'pages':len(report['pages']),
                      'characters':sum(p['characters'] for p in report['pages'])},ensure_ascii=False))
if __name__=='__main__':main()
